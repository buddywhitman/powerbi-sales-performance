import json, os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DL = os.path.join(ROOT, "deliverables")
T = json.load(open(os.path.join(DL, "test_results.json")))
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

def new(title, sub):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Segoe UI"; st.font.size = Pt(10)
    h = d.add_heading(title, 0)
    for r in h.runs: r.font.color.rgb = NAVY
    d.add_paragraph(sub)
    return d

def table(d, header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header): t.rows[0].cells[i].text = h
    for r in rows:
        cs = t.add_row().cells
        for i, v in enumerate(r): cs[i].text = str(v)
    d.add_paragraph()

def bullets(d, items):
    for i in items: d.add_paragraph(i, style="List Bullet")

# ====================================================== DESIGN DOCUMENTATION
d = new("Sales Performance Dashboard - Design Documentation", "Power BI project: SalesPerformance.pbip  |  Source: dashboard.xlsx (500 orders, 2022-2024, 27 columns)")
d.add_heading("1. Target audience analysis", 1)
table(d, ["Audience", "Needs", "Where served"], [
    ["Executive leadership", "Headline KPIs, attainment vs target, mobile access", "Sales Overview (KPI band, variance waterfall); mobile layout; paginated monthly report"],
    ["Sales managers / regional leads", "Where targets are missed; reps and regions to coach", "Geographic Performance, Customer Analysis (rep table with drill-through), Sales Rep Detail"],
    ["Product / category owners", "Mix, margin, discount impact, drivers", "Product Analysis (field-parameter explorer, treemap, decomposition tree)"],
    ["Finance / planning", "What-if on price, volume, cost", "Scenario Planner"],
    ["Distribution recipients", "Formal printable monthly pack", "Paginated report (RDL)"]])
d.add_paragraph("Adoption barrier from the brief (complex past tools) drove three rules: one consistent page template, at most five slicers per page in a fixed band, and plain-language chart titles that state the question each visual answers.")
d.add_heading("2. Page architecture and information hierarchy", 1)
table(d, ["Page", "Purpose", "Hierarchy (top to bottom)"], [
    ["Sales Overview", "How are we doing against target and budget?", "Header + nav -> 5 slicers -> KPI band -> trend (drill Year>Quarter>Month) and category chart (drill Category>Subcategory) -> variance waterfall"],
    ["Product Analysis", "What sells, at what margin?", "Slicers incl. two field parameters -> KPIs -> metric/view explorer -> category mix -> decomposition tree"],
    ["Customer Analysis", "Who buys and what drives hitting target?", "Slicers -> KPIs -> segment donut, stage funnel, rep scatter -> key influencers (AI) + rep table (drill-through)"],
    ["Geographic Performance", "Where do we win and lag?", "Slicers -> KPIs -> country/region treemap, region>country drill bars, region x category matrix"],
    ["Scenario Planner", "What if price, volume or cost change?", "3 what-if slicers + context -> scenario KPIs -> baseline vs scenario charts -> attainment gauge"],
    ["Sales Rep Detail (drill-through)", "Everything about one rep", "Back button -> KPIs -> trend + category mix -> order table"]])
d.add_heading("3. Visual branding", 1)
bullets(d, ["Palette: navy #1F3A5F (primary/actual), teal #2A9D8F (selection, scenario), orange #F4A261 (target), slate #6C7A89 (budget/baseline), purple/sky/amber/red for additional series. Page background #F4F6F9, cards white with a 6px-radius border.",
            "Typography: Segoe UI throughout (safe across Desktop, Service and mobile); Segoe UI Semibold for titles; 24pt callouts, 12pt visual titles, 10pt slicer labels.",
            "Layout standard: 1280x720, 24px outer margin, 12px gutters, 64px navy header band, 76px slicer band at y=70, KPI band at y=152, content from y=252. Implemented in a deterministic generator (build/gen_report.py) so pages cannot drift.",
            "Theme file SalesTheme-7c1d2e3f.json carries data colours, text classes and the default card style."])
d.add_heading("4. Visualization selections and rationale", 1)
table(d, ["Visual", "Where", "Why"], [
    ["Multi-value card", "KPI band on every page", "Five headline numbers in one container, consistent across pages"],
    ["Line chart (3 measures)", "Overview", "Sales vs target vs budget over time; drill Year>Quarter>Month"],
    ["Clustered column", "Overview, Scenario", "Compare actual vs target across few categories; drill Category>Subcategory"],
    ["Waterfall (advanced)", "Overview", "Shows how each region contributes to the total miss against target"],
    ["Clustered bar (field parameters)", "Product", "One visual, 6 metrics x 6 dimensions"],
    ["Treemap", "Product, Geography", "Part-to-whole with many members"],
    ["Decomposition tree (advanced)", "Product", "Ad-hoc root-cause exploration of gross profit"],
    ["Key influencers (AI)", "Customer", "Automatically ranks what makes an order On Track (>=75% of target) vs At Risk"],
    ["Donut / funnel / scatter", "Customer", "Segment share; pipeline stages; rep satisfaction vs margin with sales as size"],
    ["Table / matrix", "Customer, Geography, Rep Detail", "Exact values; region x category cross-tab; drill-through source"],
    ["Gauge", "Scenario", "Single attainment number against 100%"]])
d.add_paragraph("Total distinct visual types used: card, line, clustered column, column, waterfall, clustered bar, treemap, decomposition tree, key influencers, donut, funnel, scatter, table, matrix, gauge, slicer, button, textbox (18).")
d.add_paragraph("Known limitation: the Azure Maps visual is blocked in this tenant (admin switch for cross-region processing), so the geographic page uses a country/region treemap instead of a map. Re-enable the map by asking the tenant admin to allow Azure Maps, then swap the visual.")
d.add_heading("5. Interaction design strategy", 1)
bullets(d, ["Slicers: 25 slicer instances across pages (dropdown, 76px band). Year and Region are sync-grouped across pages; the rest are page-local to avoid surprising hidden filters.",
            "Drill-down hierarchies: Date Hierarchy (Year>Quarter>Month) on the trend line; Product Hierarchy (Category>Subcategory) on the category chart; Geography Hierarchy (Region>Country) on the geography bars. The model also defines a Customer Hierarchy.",
            "Drill-through: Sales Rep Detail filtered by Sales Rep. Right-click a rep in the Customer page table (or the scatter) and choose Drill through. A Back button returns to the source.",
            "Navigation: five page buttons in every header; the active page is highlighted teal.",
            "Bookmarks (3) on Sales Overview with buttons: Chart view and Table view swap the trend chart and a detail table (visibility only; slicers are preserved); Reset filters restores the default chart view with data state cleared.",
            "Cross-visual interactions: default cross-filtering is kept; every visual title states what it shows so highlighted results stay explainable."])
d.add_heading("6. Parameter decisions", 1)
table(d, ["Parameter", "Type", "Options", "Used by"], [
    ["Metric Selector", "Field parameter", "Sales, Gross Profit, Orders, Units Sold, Avg Order Value, Target Attainment", "Product explorer bar chart"],
    ["View By", "Field parameter", "Region, Product Category, Customer Segment, Channel, Age Group, Sales Stage", "Product explorer bar chart"],
    ["Price Change", "What-if", "-15% to +15%, step 1%", "Scenario Sales"],
    ["Volume Change", "What-if", "-25% to +25%, step 1%", "Scenario Sales / COGS"],
    ["Cost Inflation", "What-if", "0% to +20%, step 1%", "Scenario COGS"]])
d.add_paragraph("Dynamic titles are DAX measures: Title Explorer (follows both field parameters), Title Sales Trend (follows the Year filter), Title Scenario, Rep Detail Title, Title Top Performer and Selection Context. Scenario formulas: Sales x (1+price) x (1+volume); COGS x (1+volume) x (1+inflation); profit and attainment derive from those.")
d.add_paragraph("Data note: no order in the source reaches its Target (Sales < Target on all 500 rows; overall attainment 76.9%), so a binary Met/Missed flag carries no information. Target Status is therefore defined as On Track when Sales >= 75% of Target (282 orders) and At Risk otherwise (218). Change the 0.75 threshold in the Sales table's Power Query step to retune it.")
d.add_heading("7. Mobile design considerations", 1)
d.add_paragraph("A phone layout is defined for Sales Overview (visual-level mobile.json files, 360px wide): KPI card first at full width, then the trend line, then the category chart. Rationale: executives scan the headline, then the trend, then the breakdown; the dense waterfall, slicers and bookmark buttons are omitted from the phone canvas (navigation buttons on the desktop pages are 44px tall). Status: authored and schema-validated; not previewed on a phone or in the Power BI mobile app in this session.")
d.add_heading("8. Model summary", 1)
bullets(d, ["Tables: Sales (fact, loaded from the Data sheet via Power Query, 3 derived columns), Date (generated 2022-2024), Metric Selector, View By, Price Change, Volume Change, Cost Inflation.",
            "One relationship: Sales[Order Date] -> Date[Date] (many-to-one, single direction). Measures are grouped in display folders: Core, Performance, Time Intelligence, Dynamic Text, Scenario (41 total).",
            "Source path is the parameter SourcePath (expressions.tmdl); update it if dashboard.xlsx moves."])
d.add_heading("9. Performance optimization", 1)
bullets(d, ["Import mode, star-like model with one relationship; all measures are simple aggregations or iterate over at most a few dozen rows.",
            "Only needed columns are exposed; source helper columns (Year, Quarter, Month, MonthNumber, Date) are hidden in favour of the Date table.",
            "Page visual counts: 17-21 visuals per page including layout shapes and buttons; no more than 9 data visuals on any page.",
            "No bidirectional relationships, no calculated columns on the fact table (derived columns are computed in Power Query)."])
d.add_heading("10. Screenshots", 1)
for n in ["Sales Overview", "Product Analysis", "Customer Analysis", "Geographic Performance", "Scenario Planner", "Sales Rep Detail"]:
    d.add_heading(n, 3)
    d.add_picture(os.path.join(DL, "cropped_" + n.replace(" ", "_") + ".png"), width=Inches(6.3))
d.save(os.path.join(DL, "Design_Documentation.docx"))

# ====================================================== TEST REPORT
d = new("Testing Validation Report", "Scope: semantic model, rendering, parameters, interactions, performance. Environment: Power BI Desktop 2.158.1304.0 (x64), Windows 11.")
d.add_heading("1. Automated model tests (build/run_tests.py, live Desktop model)", 1)
table(d, ["Area", "Test", "Expected", "Actual", "Result"], [[r["area"], r["test"], r["expected"] if not isinstance(r["expected"], float) else round(r["expected"], 2), (round(r["actual"], 2) if isinstance(r["actual"], float) else str(r["actual"])[:70]), "PASS" if r["pass"] else "FAIL"] for r in T["results"]])
d.add_paragraph("Result: %d of %d passed. Totals reconcile to the Excel source (Total Sales $1,247,941.61; Target $1,623,770.12; 500 orders)." % (sum(r["pass"] for r in T["results"]), len(T["results"])))
d.add_paragraph("Test-authoring note: one test initially failed because it matched a what-if value by exact float equality (GENERATESERIES yields 0.1000000000000003); the model was correct and the test now matches within a tolerance.")
d.add_heading("2. Static validation", 1)
bullets(d, ["powerbi-report-author validate on SalesPerformance.Report: succeeded, 0 errors, 0 warnings (after fixing unknown formatting properties and slicer/textbox heights it reported).",
            "TMDL loaded into Power BI Desktop and the Power BI modeling MCP: 7 tables, 41 measures, 1 relationship."])
d.add_heading("3. Rendering review (Desktop screenshots, all 6 pages)", 1)
table(d, ["Finding", "Resolution"], [
    ["Nav and bookmark button labels clipped", "Added explicit zero padding and 44px height; re-captured, labels fully visible"],
    ["Field-parameter chart showed all six measures at once", "Parameter slicers given a single-select default (Sales, Region); chart now shows one metric by one dimension"],
    ["KPI values abbreviated ($1M, $2K)", "Display units set to none; values now exact"],
    ["Key influencers: 'not enough data'", "Cause: all 500 orders missed target so one class had no members. Re-defined Target Status (On Track >= 75% of target); visual now analyses"],
    ["Azure Maps blocked by tenant switch", "Replaced with treemap; limitation documented"],
    ["Decomposition tree rendered blank at 222px", "Enlarged to 600x452; root node and expand control render"]])
d.add_heading("4. Functional areas and status", 1)
table(d, ["Area", "What was verified", "Not verified"], [
    ["Functionality", "Measures reconcile to source; every page renders data; dynamic title measures return text", "-"],
    ["Cross-filtering", "Region, Year and Category filter contexts return source-matching totals (DAX level)", "Clicking visuals in Desktop to confirm highlight behaviour"],
    ["Parameters", "Field parameters have 6 fields each and drive the explorer (default selection renders correctly); what-if math verified for +10% price; scenario equals baseline at defaults", "Moving the what-if slicers interactively in the UI"],
    ["Drill-through", "Page binding, filter and Back button validate; every rep key resolves (7 reps)", "Right-click drill-through gesture in Desktop"],
    ["Bookmarks", "3 bookmark files, index and button targets pass validator checks", "Clicking each button and the reverse transitions"],
    ["Mobile", "Mobile layout files validate", "Phone preview / mobile app rendering"],
    ["Paginated report", "RDL is well-formed XML; the same DAX datasets run against the model (Year=2024: 193 orders, $476,169 sales); sample PDF produced", "Opening the RDL in Power BI Report Builder or the service; the sample PDF is a reportlab rendering of the same queries, not a Report Builder export"]])
d.add_heading("5. Performance", 1)
d.add_paragraph("Heaviest page-style query (Year x Quarter x Region x Category with four measures), five runs, wall-clock through the pbi CLI: %s ms; median %s ms. These timings include CLI process start-up (about 1.4 s on trivial queries), so engine time is well below these figures. They are an upper bound, not a benchmark. Dataset is 500 rows." % (", ".join(str(x) for x in T["perf"]["heavy_query_ms"]), T["perf"]["median_ms"]))
d.add_paragraph("Responsive behaviour: pages are authored at 1280x720 with FitToPage; screenshots were taken in a single window size only.")
d.save(os.path.join(DL, "Testing_Validation_Report.docx"))

# ====================================================== USER GUIDE
d = new("Sales Performance Dashboard - Quick Guide", "For executives, sales managers and analysts")
for t, b in [
    ("Moving around", ["Use the five buttons in the top-right of every page. The highlighted (teal) button is the page you are on.", "Year and Region selections carry across pages that have those slicers."]),
    ("Filtering", ["Each page has a row of dropdown slicers under the header. Choose one or more values; clear with the eraser that appears on hover. 'All' means no filter."]),
    ("Drilling down", ["Sales Overview: hover the trend chart and use the drill arrows (down-arrow icon at the top-right of the visual) to go Year > Quarter > Month.", "Product chart on Sales Overview: drill from category to subcategory. Geography bars: region to country."]),
    ("Drill-through to a sales rep", ["On Customer Analysis, right-click a rep in the table (or a dot in the scatter) > Drill through > Sales Rep Detail. Use '< Back' to return."]),
    ("Switching the Overview view", ["'Chart view' and 'Table view' swap the trend chart for a table of the same data; 'Reset filters' restores the default."]),
    ("Choosing what Product Analysis shows", ["Pick a metric and a dimension in the two 'field parameter' slicers; the chart title updates to say what you are looking at."]),
    ("Running scenarios", ["On Scenario Planner pick a price, volume and cost-inflation change. The banner under the slicers restates the scenario; the KPI band shows the sales/profit impact against the baseline."]),
    ("Reading Target Status", ["On Track = sales at least 75% of the order's target; At Risk = below that. Key influencers lists which attributes raise the odds of being At Risk or On Track."]),
    ("Printing the monthly pack", ["Run SalesPerformance_Monthly.rdl in Power BI Report Builder or the Power BI service; choose Year, Region, Product category and the recipient name, then export to PDF."])]:
    d.add_heading(t, 2)
    bullets(d, b)
d.save(os.path.join(DL, "User_Guide.docx"))
print("docs written")
